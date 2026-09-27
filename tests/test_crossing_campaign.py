"""Source census and checkpoint behaviour for crossing-change exploration."""
import hashlib
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path
import unittest
from unittest.mock import patch

import netCDF4
import numpy as np
from openpyxl import Workbook
from test_coxeter import S5_GAUSS, S5_SEEDS
from test_crossing_change import SIX

SCRIPT = Path(__file__).resolve().parents[1] / 'scripts' / '11_crossing_changes.py'


class CrossingCampaignTests(unittest.TestCase):
    def test_cli_checkpoints_a_source_and_resume_skips_it(self):
        known = SCRIPT.parents[1] / 'results' / 'ten_new_stick_knots' / 'K15n59007_equilateral_10sticks.txt'
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'eddy').mkdir()
            (root / 'results').mkdir()
            (root / 'eddy' / known.name).write_bytes(known.read_bytes())
            with netCDF4.Dataset(root / 'empty.nc', 'w'):
                pass
            wb = Workbook()
            wb.active.append(['15n41189', S5_GAUSS, 4, None, 1, S5_SEEDS, 0, None, 0, None])
            wb.save(root / 'rows.xlsx')
            args = [sys.executable, str(SCRIPT), '--eddy', str(root / 'eddy'),
                    '--results', str(root / 'results'), '--workbook', str(root / 'rows.xlsx'),
                    '--output', str(root / 'campaign.json'), '--scratch', str(root / 'scratch'),
                    '--limit', '1', '--source-budget', '30', '--candidate-budget', '2',
                    '--strategy', 'uniform', '--fractions', '.5', '--weights', '.25',
                    '--times', '.5']
            env = {**os.environ, 'EQUISTICK_CRSS': str(root / 'empty.nc')}
            first = subprocess.run(args, env=env, capture_output=True, text=True,
                                   timeout=45, check=True)
            report = json.loads((root / 'campaign.json').read_text())
            self.assertEqual(report['inventory_count'], 1)
            self.assertEqual(report['completed_count'], 1)
            self.assertEqual(report['sources']['eddy:K15n59007_equilateral_10sticks']['moves_tried'], 140)
            second = subprocess.run(args, env=env, capture_output=True, text=True,
                                    timeout=45, check=True)
            self.assertEqual(json.loads((root / 'campaign.json').read_text())['completed_count'], 1)
            self.assertIn('skipped_completed=1', second.stdout)

    def test_invalid_crossing_witness_cannot_publish_a_polygon(self):
        spec = importlib.util.spec_from_file_location('crossing_campaign', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        path = SCRIPT.parents[1] / 'results' / 'ten_new_stick_knots' / 'K15n59007_equilateral_10sticks.txt'
        source = {'id': 'ours:K15n59007', 'label': 'K15n59007', 'kind': 'ours',
                  'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
        V = np.loadtxt(path)
        witness = {'coordinates': V.tolist(), 'move': {'vertex': 0,
                   'destination': V[0].tolist()}, 'crossing': {'edge': 3}}
        with tempfile.TemporaryDirectory() as directory:
            result = module.attempt_equalization(source, witness, 'K15n59007',
                                                 Path(directory), budget_seconds=2)
            self.assertEqual(result['status'], 'invalid_crossing')
            self.assertEqual(list(Path(directory).glob('*.txt')), [])

    def test_final_saved_file_gets_interval_geometry_and_three_fresh_type_checks(self):
        spec = importlib.util.spec_from_file_location('crossing_campaign', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        path = SCRIPT.parents[1] / 'results' / 'ten_new_stick_knots' / 'K15n59007_equilateral_10sticks.txt'
        checks = module.validate_saved_candidate(path, 'K15n59007')
        self.assertEqual(checks['interval']['status'], 'certified')
        self.assertTrue(checks['high_precision'])
        self.assertLess(checks['mr_ratio'], 1)
        self.assertEqual(checks['identity']['name'], 'K15n59007')
        self.assertEqual(len(checks['identity']['projection_seeds']), 3)
        self.assertEqual(checks['sha256'], hashlib.sha256(path.read_bytes()).hexdigest())

    def test_summary_logs_every_wirt_four_hit_but_only_queues_exact_maps(self):
        spec = importlib.util.spec_from_file_location('crossing_campaign', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        with tempfile.TemporaryDirectory() as directory:
            workbook = Path(directory) / 'rows.xlsx'
            wb = Workbook()
            wb.active.append(['15n41189', S5_GAUSS, 4, None, 1, S5_SEEDS, 0, None, 0, None])
            wb.active.append(['15n59060', S5_GAUSS, 4, None, 0, None, 0, None, 0, None])
            wb.save(workbook)
            index = module.build_wirt_index(workbook)
            source_result = {'status': 'completed', 'source': {'id': 'eddy:Source',
                             'label': 'Source', 'kind': 'eddy', 'path': '/example'},
                             'source_sha256': 'abcdef', 'elapsed_seconds': 5,
                             'scan': {'moves_tried': 840, 'single_crossings': 4,
                                      'identified_counts': {'K15n41189': 2, 'K15n59060': 1,
                                                            'unknot': 1},
                                      'inconclusive_count': 0, 'candidates': {
                                          'K15n41189': [{'move_index': 1}],
                                          'K15n59060': [{'move_index': 2}]}}}
            summary, targets = module.summarize_source(source_result, index)
            self.assertEqual(summary['moves_tried'], 840)
            self.assertEqual(set(summary['four_bridge_reached']),
                             {'K15n41189', 'K15n59060'})
            self.assertTrue(summary['four_bridge_reached']['K15n41189']['exact_rank_four_passed'])
            self.assertFalse(summary['four_bridge_reached']['K15n59060']['exact_rank_four_passed'])
            self.assertEqual(set(targets), {'K15n41189'})

    def test_source_record_includes_checked_hash_identity_and_all_move_attempts(self):
        spec = importlib.util.spec_from_file_location('crossing_campaign', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        path = SCRIPT.parents[1] / 'results' / 'ten_new_stick_knots' / 'K15n59007_equilateral_10sticks.txt'
        source = {'id': 'ours:K15n59007', 'label': 'K15n59007', 'kind': 'ours',
                  'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
        result = module.process_source(source, fractions=(.5,), weights=(.25,),
                                       times=(.5,), strategy='uniform')
        self.assertEqual(result['status'], 'completed')
        self.assertEqual(result['source']['id'], source['id'])
        self.assertEqual(result['source_sha256'], source['sha256'])
        self.assertEqual(result['scan']['moves_tried'], 140)
        self.assertEqual(result['scan']['source_identification']['name'], source['label'])

    def test_new_targets_need_wirtinger_four_and_exact_coxeter_map(self):
        spec = importlib.util.spec_from_file_location('crossing_campaign', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'rows.xlsx'
            wb = Workbook()
            wb.active.append(['15n41189', S5_GAUSS, 4, None, 1, S5_SEEDS, 0, None, 0, None])
            wb.active.append(['15n59060', S5_GAUSS, 4, None, 0, None, 0, None, 0, None])
            wb.active.append(['15n00001', S5_GAUSS, 3, None, 1, S5_SEEDS, 0, None, 0, None])
            wb.save(path)
            index = module.build_wirt_index(path)
            valid = module.classify_reached('K15n41189', index)
            self.assertTrue(valid['exact_rank_four_passed'])
            self.assertEqual(valid['exact_map']['relations_checked'], 15)
            self.assertEqual(valid['exact_map']['group_order'], 120)
            missing = module.classify_reached('K15n59060', index)
            self.assertTrue(missing['wirtinger_four'])
            self.assertFalse(missing['exact_rank_four_passed'])
            self.assertFalse(module.classify_reached('K15n00001', index)['wirtinger_four'])

    def test_rolfsen_name_wins_over_ht_alias_for_same_eight_crossing_knot(self):
        spec = importlib.util.spec_from_file_location('crossing_campaign', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        names = ['o9_43592', '8_13', 'K9_744', 'K8a7']
        self.assertEqual(module.select_table_name(names), '8_13')

    def test_scanner_counts_every_move_and_records_only_single_crossing_endpoints(self):
        spec = importlib.util.spec_from_file_location('crossing_campaign', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        report = module.scan_polygon(SIX, fractions=(.5,), weights=(.25,), times=(.5,),
                                     strategy='uniform')
        self.assertEqual(report['moves_tried'], 36)
        self.assertGreater(report['single_crossings'], 0)
        self.assertEqual(report['single_crossings'],
                         sum(report['identified_counts'].values()) + report['inconclusive_count'])
        for witnesses in report['candidates'].values():
            for witness in witnesses:
                move = witness['move']
                crossing = module.find_single_crossing(SIX, move['vertex'],
                                                       np.asarray(move['destination']))
                self.assertIsNotNone(crossing)
                self.assertEqual(crossing['edge'], witness['crossing']['edge'])

    def test_nearest_grid_is_included_by_default(self):
        spec = importlib.util.spec_from_file_location('crossing_campaign', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        report = module.scan_polygon(SIX, strategy='nearest')
        self.assertEqual(report['moves_tried'], 648)
        self.assertGreater(report['single_crossings'], 0)

    def test_saved_ten_gon_is_identified_with_two_meridian_preserving_projections(self):
        spec = importlib.util.spec_from_file_location('crossing_campaign', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        path = SCRIPT.parents[1] / 'results' / 'ten_new_stick_knots' / 'K15n59007_equilateral_10sticks.txt'
        proof = module.identify_polygon(np.loadtxt(path), needed=2)
        self.assertEqual(proof['status'], 'identified', proof)
        self.assertEqual(proof['name'], 'K15n59007')
        self.assertEqual(len(proof['projection_seeds']), 2)

    def test_inventory_includes_all_three_sources_and_only_certified_local_ten_gons(self):
        spec = importlib.util.spec_from_file_location('crossing_campaign', SCRIPT)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            eddy = root / 'mseq_knots'
            results = root / 'results'
            eddy.mkdir()
            results.mkdir()
            coords = np.arange(30, dtype=float).reshape(10, 3)
            np.savetxt(eddy / '8_19.txt', coords)
            np.savetxt(eddy / 'TooLong.txt', np.arange(33).reshape(11, 3))
            local = results / 'K13n225_equilateral_10sticks.txt'
            np.savetxt(local, coords)
            (results / 'Uncertified_equilateral_10sticks.txt').write_text('not a polygon')
            (results / 'interval_certificates.json').write_text(json.dumps({'files': [
                {'file': local.name, 'status': 'certified', 'sticks': 10,
                 'sha256': hashlib.sha256(local.read_bytes()).hexdigest()}]}))
            nc = root / 'fixture.nc'
            with netCDF4.Dataset(nc, 'w') as ds:
                good = ds.createGroup('K13n586')
                good.setncattr('sticks', 10)
                good.setncattr('crossings', 13)
                short = ds.createGroup('Other')
                short.setncattr('sticks', 9)
                short.setncattr('crossings', 10)
            with patch.dict('os.environ', {'EQUISTICK_CRSS': str(nc)}):
                sources = module.collect_sources(eddy, results)
            self.assertEqual([source['id'] for source in sources],
                             ['crss:K13n586', 'eddy:8_19', 'ours:K13n225'])


if __name__ == '__main__':
    unittest.main()
